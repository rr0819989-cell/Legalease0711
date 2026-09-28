import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def _clean_heading(line: str):
    return re.sub(r"^#+\s*", "", line).strip()


def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Header / simple LegalEase branding
    header = section.header.paragraphs[0]
    header.text = "⚖ LegalEase | AI Legal Document Generator"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Footer
    footer = section.footer.paragraphs[0]
    footer.text = "Generated using LegalEase. Review the draft before signing."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type)
    run.bold = True
    run.font.size = Pt(18)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            continue

        if line.startswith("#"):
            p = doc.add_paragraph()
            r = p.add_run(_clean_heading(line))
            r.bold = True
            r.font.size = Pt(13 if line.startswith("##") else 15)
        elif re.match(r"^\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 11)
        self.cell(0, 8, "LegalEase", ln=1, align="C")
        self.set_draw_color(120, 120, 120)
        self.line(10, 18, 200, 18)
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, "LegalEase - Review before signing | Page " + str(self.page_no()), align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, sanitize_text(doc_type), align="C")
    pdf.ln(3)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            pdf.ln(3)
            continue

        if line.startswith("#"):
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 7, _clean_heading(line))
            pdf.ln(1)
        elif re.match(r"^\d+\.\s+", line):
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, "• " + re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, "• " + line[2:])
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 6, line)

    return bytes(pdf.output())


def format_html_preview(text: str) -> str:
    html_lines = []
    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            html_lines.append("<br>")
        elif line.startswith("#"):
            html_lines.append(f"<h3>{_clean_heading(line)}</h3>")
        elif re.match(r"^\d+\.\s+", line):
            html_lines.append(f"<li>{re.sub(r'^\d+\.\s+', '', line)}</li>")
        elif line.startswith("- "):
            html_lines.append(f"<li>{line[2:]}</li>")
        else:
            html_lines.append(f"<p>{line}</p>")
    return "\n".join(html_lines)
